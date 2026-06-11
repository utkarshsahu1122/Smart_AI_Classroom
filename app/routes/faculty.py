import os
import time
import uuid
from io import BytesIO
from datetime import datetime
from flask import Blueprint, request, jsonify, current_app, send_file, Response
from werkzeug.utils import secure_filename
from ..services.timezone import to_local
from ..models.schemas import (
    create_quiz_session, get_quiz_session_by_code, insert_quiz_questions,
    get_questions_by_session, get_students_by_session, get_assigned_questions,
    update_quiz_session, update_student, count_assigned_questions,
)
from ..services.qr_service import generate_qr_for_session

bp = Blueprint('faculty', __name__, url_prefix='/api/faculty')

ALLOWED_EXTENSIONS = {'pdf'}

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@bp.route('/upload', methods=['POST'])
def upload_pdf():
    """Upload a PDF, generate vectorstore + question pool, and create a session."""
    start_time = time.time()
    print(f"Request started at: {start_time}")
    
    if 'pdf_file' not in request.files:
        return jsonify({'success': False, 'error': 'No file part'}), 400
    file = request.files['pdf_file']
    if file.filename == '':
        return jsonify({'success': False, 'error': 'No selected file'}), 400
    if not file or not allowed_file(file.filename):
        return jsonify({'success': False, 'error': 'Only PDF files are allowed'}), 400

    filename = secure_filename(file.filename)
    mcq_weight = int(request.form.get('mcq_weight', 70))
    fill_weight = int(request.form.get('fill_weight', 30))
    
    session_code = str(uuid.uuid4())[:8].upper()
    uploads_dir = os.path.join(current_app.root_path, 'static', 'temp_uploads')
    os.makedirs(uploads_dir, exist_ok=True)

    unique_filename = f"{session_code}_{filename}"
    file_path = os.path.join(uploads_dir, unique_filename)
    file.save(file_path)

    # Process PDF + generate entire question pool at upload time
    from ..services.rag_engine import process_pdf_and_generate_pool
    try:
        quiz_data = process_pdf_and_generate_pool(file_path, session_code)
    except Exception as e:
        print(f"[Faculty] Error during PDF processing or generation: {e}")
        try:
            os.remove(file_path)
        except OSError:
            pass
        end_time = time.time()
        print(f"Request ended at: {end_time}. Duration: {end_time - start_time:.2f}s")
        return jsonify({'success': False, 'error': 'Processing exceeded limits or failed. Please try a smaller PDF or try again later.'}), 500

    # VALIDATION: Never create a session with 0 questions
    if not quiz_data or len(quiz_data) < 5:
        try:
            os.remove(file_path)
        except OSError:
            pass
        count = len(quiz_data) if quiz_data else 0
        print(f"[Faculty] REJECTED session {session_code}: only {count} questions generated")
        return jsonify({
            'success': False,
            'error': f'Question generation failed — only {count} questions could be generated. This can happen if the PDF is too short, image-heavy, or the AI service is temporarily unavailable. Please try uploading again.'
        }), 422

    qr_path = generate_qr_for_session(session_code)

    # Create session document in MongoDB
    session_doc = create_quiz_session(
        session_code=session_code,
        pdf_filename=unique_filename,
        qr_code_path=qr_path,
        timer_minutes=10,
        mcq_weight=mcq_weight,
        fill_weight=fill_weight
    )

    # Insert questions into MongoDB
    insert_quiz_questions(session_doc["id"], quiz_data)
    print(f"[Faculty] Session {session_code} created with {len(quiz_data)} questions in pool.")

    end_time = time.time()
    print(f"Request ended at: {end_time}. Duration: {end_time - start_time:.2f}s")

    return jsonify({
        'success': True,
        'session_code': session_code,
        'qr_url': f"/static/{qr_path}" if qr_path else None,
        'questions_count': len(quiz_data),
        'pdf_filename': unique_filename
    })

@bp.route('/session/<session_code>')
def session_details(session_code):
    """Get session details including student list."""
    quiz_session = get_quiz_session_by_code(session_code)
    if not quiz_session:
        return jsonify({'success': False, 'error': 'Session not found'}), 404

    students = get_students_by_session(quiz_session["id"])
    questions = get_questions_by_session(quiz_session["id"])

    student_list = []
    for s in students:
        assigned_count = count_assigned_questions(s["id"])
        student_list.append({
            'id': s['id'],
            'roll_no': s['roll_no'],
            'name': s['name'],
            'score': s.get('score', 0),
            'total': assigned_count,
            'is_logged_in': s.get('is_logged_in', False),
            'submitted': s.get('submitted_at') is not None,
            'submitted_at': to_local(s['submitted_at']) if s.get('submitted_at') else None,
            'unfair_means': s.get('unfair_means', False),
            'warning_count': s.get('warning_count', 0),
        })

    return jsonify({
        'success': True,
        'session': {
            'session_code': quiz_session['session_code'],
            'pdf_filename': quiz_session['pdf_filename'],
            'qr_url': f"/static/{quiz_session['qr_code_path']}" if quiz_session.get('qr_code_path') else None,
            'created_at': to_local(quiz_session['created_at']),
            'timer_minutes': quiz_session.get('timer_minutes', 10),
            'is_active': quiz_session.get('is_active', True),
            'questions_count': len(questions),
        },
        'students': student_list,
    })

@bp.route('/session/<session_code>/end', methods=['POST'])
def end_session(session_code):
    """Faculty ends the session — auto-submits all pending students."""
    quiz_session = get_quiz_session_by_code(session_code)
    if not quiz_session:
        return jsonify({'success': False, 'error': 'Session not found'}), 404

    update_quiz_session(quiz_session["id"], {
        "is_active": False,
        "ended_at": datetime.utcnow(),
    })

    students = get_students_by_session(quiz_session["id"])
    auto_count = 0
    for student in students:
        if student.get("submitted_at") is None:
            assigned = get_assigned_questions(student["id"])
            score = 0
            for sq in assigned:
                if sq.get("student_answer"):
                    if sq["student_answer"].strip().lower() == sq["question"]["correct_answer"].strip().lower():
                        score += 1
            update_student(student["id"], {
                "score": score,
                "submitted_at": datetime.utcnow(),
                "is_logged_in": False,
            })
            auto_count += 1

    print(f"[Faculty] Session {session_code} ended. {auto_count} students auto-submitted.")

    return jsonify({
        'success': True,
        'auto_submitted': auto_count
    })

@bp.route('/session/<session_code>/report')
def download_report(session_code):
    """Generate and download a .docx Word report for the quiz session."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    quiz_session = get_quiz_session_by_code(session_code)
    if not quiz_session:
        return jsonify({'success': False, 'error': 'Session not found'}), 404

    students = get_students_by_session(quiz_session["id"])

    doc = Document()

    title = doc.add_heading('Quiz Session Report', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading('Session Information', level=2)
    info_table = doc.add_table(rows=5, cols=2, style='Light Shading Accent 1')
    info_data = [
        ('Session Code', quiz_session['session_code']),
        ('PDF Document', quiz_session['pdf_filename']),
        ('Created At', to_local(quiz_session['created_at'])),
        ('Status', 'Ended' if not quiz_session.get('is_active') else 'Active'),
        ('Total Students', str(len(students))),
    ]
    for i, (label, value) in enumerate(info_data):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
        for cell in info_table.rows[i].cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)

    doc.add_paragraph('')

    doc.add_heading('Student Results', level=2)

    if students:
        result_table = doc.add_table(rows=1, cols=5, style='Light Shading Accent 1')
        hdr_cells = result_table.rows[0].cells
        hdr_cells[0].text = 'S.No'
        hdr_cells[1].text = 'Roll No'
        hdr_cells[2].text = 'Name'
        hdr_cells[3].text = 'Score'
        hdr_cells[4].text = 'Submitted At'

        for idx, student in enumerate(students, 1):
            row = result_table.add_row()
            total_q = count_assigned_questions(student["id"]) or 20
            submitted = to_local(student['submitted_at']) if student.get('submitted_at') else 'Not Submitted'
            values = [str(idx), student['roll_no'], student['name'], f"{student.get('score', 0)}/{total_q}", submitted]
            for i, val in enumerate(values):
                cell = row.cells[i]
                cell.text = val
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
    else:
        doc.add_paragraph('No students participated in this session.')

    doc.add_paragraph('')

    # --- Summary Statistics ---
    if students:
        scores = [s.get('score', 0) for s in students if s.get('submitted_at')]
        if scores:
            doc.add_heading('Summary Statistics', level=2)
            total_q = count_assigned_questions(students[0]["id"]) or 20
            stats_para = doc.add_paragraph()
            stats_para.add_run('Average Score: ').bold = True
            stats_para.add_run(f'{sum(scores)/len(scores):.1f}/{total_q}\n')
            stats_para.add_run('Highest Score: ').bold = True
            stats_para.add_run(f'{max(scores)}/{total_q}\n')
            stats_para.add_run('Lowest Score: ').bold = True
            stats_para.add_run(f'{min(scores)}/{total_q}\n')
            
            pass_threshold = total_q / 2
            stats_para.add_run(f'Pass Rate (≥{pass_threshold}): ').bold = True
            pass_count = sum(1 for s in scores if s >= pass_threshold)
            stats_para.add_run(f'{pass_count}/{len(scores)} ({pass_count/len(scores)*100:.0f}%)')

    doc.add_paragraph('')
    footer = doc.add_paragraph(f'Report generated on {datetime.utcnow().strftime("%d-%m-%Y %H:%M")} UTC')
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(150, 150, 150)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    filename = f"Quiz_Report_{session_code}.docx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
